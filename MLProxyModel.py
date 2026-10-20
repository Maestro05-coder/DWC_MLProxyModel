"""
ML-assisted history matching and 20-year forecast
=====================================================================

HOW TO RUN
    pip install numpy pandas matplotlib openpyxl scipy scikit-learn
    python HistoryMatchingJudges.py path/to/folder      (or edit DATA_DIR below)

    DATA folder is used to store the original files from the competition

    MLProxyModel_Results folder is created to hold the results (Excel + PNGs)

"""
import datetime
import re
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.patheffects as pe
from matplotlib.patches import Patch
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.optimize import curve_fit, differential_evolution
from scipy.stats import qmc
from sklearn.decomposition import PCA
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, Matern, WhiteKernel

# =============================================================================
# CONFIG - the only part you should need to edit
# =============================================================================
DATA_DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("data")
OUTPUT_DIR = Path("MLProxyModel_Results")
SHOW_PLOTS = True                 # False = only save the PNGs
HM_XLSX = "1_History_Matching_Results.xlsx"   # part A: history matching
FC_XLSX = "2_Forecast_Results.xlsx"           # part B: 20-year forecast

FILES = {                         # file names without .xlsx
    "intro":    "01 Introduction",
    "field":    "02 Field Production",
    "template": "09 Template Deliverable",
}
PARAM_SHEET = "Uncertainty Values"
PRODUCTION_SHEETS = {
    "gas_cum":    "Gas Production Cumulative",
    "gas_rate":   "Gas Production Rate",
    "oil_cum":    "Oil production cumulative",
    "oil_rate":   "Oil Production Rate",
    "water_cum":  "Water Production Cumulative",
    "water_rate": "Water Production Rate",
}
HISTORY_COL = "RAMP_History"      # the real observed field
CASE_PREFIX = "Case_"

# The official split lives in the files 03-08 (matched by the number their file name starts with).
SPLIT_FILE_PREFIX = {"03": "train", "04": "train", "05": "train", "06": "train",
                     "07": "validation", "08": "test"}
# FALLBACK only (used if files 03-08 are missing or cannot be matched to Case_N): ASSUMED ranges
TRAIN_IDS = range(1, 71)
VAL_IDS = range(71, 86)
TEST_IDS = range(86, 101)

# Streams used to SCORE the match and how each is transformed first.
# Gas is left out (it follows oil, so it would count oil twice); oil and gas RATES are
# left out (every case starts equal to the history, so they say little).
MATCH_STREAMS = {"oil_cum": "linear", "water_cum": "linear", "water_rate": "log"}
SIM_START = None                  # None = first date found in RAMP_History
LOG_PARAMS = ["Permeability Multiplier"]   # settings handled on a log scale
PCA_VAR, PCA_MAX = 0.999, 6       # curve compression
N_SOBOL_POW2 = 14                 # 2**14 trial settings for the search
ACCEPT_MARGIN_MIN = 0.01          # a trial is "as good as the best" if its misfit is within the
ACCEPT_MARGIN_REL = 0.10          # LARGEST of: this amount, this share of the best misfit,
                                  # or 2x the proxy's own error
FINAL_REFIT_ON_VAL = True         # after reporting accuracy, refit on train+validation
RANDOM_SEED = 0

# --- forecast 2008 -> 2028 and the submission template
FORECAST_START_IF_NO_TEMPLATE = "2008-04-01"
FORECAST_END_IF_NO_TEMPLATE = "2028-01-01"
FIT_WINDOW_YEARS = None           # None = choose automatically with the hindcast
FIT_WINDOW_CANDIDATES = [2.0, 3.0, 5.0]
HINDCAST_YEARS = 3.0              # hide this many years of history to test the forecast
WOR_MAX = 50.0                    # cap on water-oil ratio (50 = 98% water cut)
ANCHOR_TO_HISTORY = True          # make forecasts join the real 2008 history exactly
FORECAST_REFIT_ON_ALL = False     # False = test cases are never used for fitting
MIN_FORECAST_SETS, MAX_FORECAST_SETS = 50, 500
N_SELECT = 10                     # cases to submit (= sheets in the template)
# "diverse_params": best match first, then the most different settings that still fit
# "best":           the lowest-misfit settings (often near duplicates)
# "span_forecast":  settings spread from low to high final oil forecast
SELECTION_MODE = "diverse_params"

PLAIN = {"oil_cum": "Oil produced so far (cumulative)",
         "water_cum": "Water produced so far (cumulative)",
         "water_rate": "Water production rate"}
UNITS = {"oil_cum": "STB", "water_cum": "STB", "water_rate": "STB/day",
         "gas": "MSCF", "oil": "STB", "water": "STB"}
STREAMS = ["gas", "oil", "water"]                  # order of the template columns
GREEN, AMBER, RED, BLUE = "#1a7f37", "#b8860b", "#b22222", "#1f77b4"


# =============================================================================
# STEP 1 - FIND THE EXCEL FILES AND READ THE SHEETS
# =============================================================================
class SheetNotFound(Exception):
    pass


_NAMES_CACHE = {}


def listWorkbooks():
    return [p for p in sorted(DATA_DIR.glob("*.xlsx")) if not p.name.startswith("~$")]


def findFile(stem):
    """The .xlsx in DATA_DIR whose name matches `stem` (exact, then 'starts with')."""
    target = stem.strip().lower()
    for p in listWorkbooks():
        if p.stem.strip().lower() == target:
            return p
    for p in listWorkbooks():
        if p.stem.strip().lower().startswith(target):
            return p
    return None


def sheetNames(path):
    """Sheet names of a workbook without loading its cells (fast)."""
    key = str(path)
    if key not in _NAMES_CACHE:
        with pd.ExcelFile(path) as xl:
            _NAMES_CACHE[key] = xl.sheet_names
    return _NAMES_CACHE[key]


def findWorkbookWithSheet(sheetName):
    """Fallback for ONE combined workbook: the first .xlsx that has this sheet."""
    target = sheetName.strip().lower()
    for p in listWorkbooks():
        if target in [n.strip().lower() for n in sheetNames(p)]:
            return p
    return None


def readSheet(fileStem, sheetName, header=None):
    path = findFile(fileStem)
    if path is None:
        path = findWorkbookWithSheet(sheetName)
    if path is None:
        raise FileNotFoundError(
            f"Could not find a file named like '{fileStem}.xlsx' in {DATA_DIR.resolve()}.\n"
            f"Excel files I can see there: {[p.name for p in listWorkbooks()]}")
    lookup = {n.strip().lower(): n for n in sheetNames(path)}
    real = lookup.get(sheetName.strip().lower())
    if real is None:
        raise SheetNotFound(f"Sheet '{sheetName}' not found in '{path.name}'. "
                            f"Sheets there: {sheetNames(path)}")
    return pd.read_excel(path, sheet_name=real, header=header)


def loadParameters():
    """Uncertainty Values -> one row per case ('Case_1'...), one column per setting."""
    df = readSheet(FILES["intro"], PARAM_SHEET, header=0)
    df = df.rename(columns={df.columns[0]: "Item"})
    df.columns = [str(c).strip() for c in df.columns]
    df["Item"] = df["Item"].astype(str).str.strip().str.replace(" ", "_", regex=False)
    return df.set_index("Item").apply(pd.to_numeric, errors="coerce").dropna(how="all")


def loadProduction(sheetName):
    """One production sheet -> clean (dates x cases) table. Tolerates the different
    header layouts by locating the row whose first cell says 'Date'."""
    raw = readSheet(FILES["field"], sheetName, header=None)
    headerRow = raw.index[raw.iloc[:, 0].astype(str).str.strip().eq("Date")][0]
    body = raw.iloc[headerRow + 2:, :].copy()                  # skip header + units rows
    body.columns = ["Date"] + [str(c).strip() for c in raw.iloc[headerRow, 1:]]
    body["Date"] = pd.to_datetime(body["Date"], errors="coerce")
    body = body.dropna(subset=["Date"]).set_index("Date").sort_index()
    body = body.apply(pd.to_numeric, errors="coerce")
    return body.loc[:, ~body.columns.isin(["nan", ""])]


def caseCols(df):
    return [c for c in df.columns if str(c).startswith(CASE_PREFIX)]


def lastValid(series):
    s = series.dropna()
    return s.iloc[-1] if len(s) else np.nan


def caseNumber(name):
    nums = re.findall(r"\d+", str(name))
    return int(nums[-1]) if nums else -1


SPLIT_MAP = {}                    # case number -> split, filled from files 03-08 by buildSplitMap()
SPLIT_SOURCE = "ASSUMED from case numbers (1-70 / 71-85 / 86-100): files 03-08 were not used"
SPLIT_REPORT = pd.DataFrame()     # what was found in each of files 03-08 (also written to Excel)


def caseIdsInWorkbook(path):
    """Case numbers contained in one train / validation / test workbook, WITHOUT loading the big
    grids: from the sheet names ('Case_12', 'Case 12', '12'), else from a column named like 'case'."""
    ids = []
    for n in sheetNames(path):
        m = re.search(r"case\D*(\d+)", n, re.I) or re.fullmatch(r"\s*(\d+)\s*", n)
        if m:
            ids.append(int(m.group(1)))
    if ids:
        return ids, "sheet names"
    try:
        df = pd.read_excel(path, sheet_name=0, usecols=lambda c: "case" in str(c).lower())
    except Exception:
        return [], "nothing found"
    if df.shape[1]:
        vals = df.iloc[:, 0].dropna().astype(str).str.extract(r"(\d+)")[0].dropna().astype(int)
        return sorted(vals.unique().tolist()), "'case' column"
    return [], "nothing found"


def buildSplitMap(caseNames):
    """Label every Case_N as train / validation / test from WHICH FILE it is in (03-08).
    Falls back to the assumed ranges, loudly, if the files cannot be matched."""
    global SPLIT_MAP, SPLIT_SOURCE, SPLIT_REPORT
    print("\n=== TRAIN / VALIDATION / TEST FILES (03-08) ===")
    found, clash, rows = {}, set(), []
    for p in listWorkbooks():
        split = SPLIT_FILE_PREFIX.get(p.stem.strip()[:2])
        if split is None:
            continue
        ids, how = caseIdsInWorkbook(p)
        rows.append({"file": p.name, "split": split, "cases found": len(set(ids)), "found via": how,
                     "first ids": ", ".join(str(i) for i in sorted(set(ids))[:5])})
        for i in set(ids):
            if i in found and found[i][0] != split:
                clash.add(i)
            found.setdefault(i, (split, p.name))
    SPLIT_REPORT = pd.DataFrame(rows)
    wanted = {caseNumber(c) for c in caseNames}
    reason = None
    if not rows:
        reason = "no files starting with 03-08 were found in the data folder"
    elif not found:
        reason = "the files 03-08 were found but no case numbers could be read from their sheet names or a 'case' column"
    elif clash:
        reason = (f"case numbers repeat across train / validation / test files (e.g. {sorted(clash)[:5]}), so each "
                  "file probably restarts its numbering and cannot be matched to Case_N in file 02")
    elif len(wanted & found.keys()) < 0.5 * len(wanted):
        reason = (f"only {len(wanted & found.keys())} of the {len(wanted)} cases in file 02 appear in files 03-08")
    if reason:
        SPLIT_MAP = {}
        SPLIT_SOURCE = f"ASSUMED from case numbers (1-70 / 71-85 / 86-100) because {reason}"
        print(SPLIT_REPORT.to_string(index=False) if len(rows) else "  (no split files)")
        print(f"  WARNING: {reason}.\n  -> falling back to the assumed ranges. Tell me how the cases are numbered "
              "in those files to fix this.")
        return
    SPLIT_MAP = {i: s for i, (s, _) in found.items()}
    SPLIT_SOURCE = "read from the files 03-08 (which file each case is in)"
    print(SPLIT_REPORT.to_string(index=False))
    counts = pd.Series([SPLIT_MAP.get(caseNumber(c), "unassigned") for c in caseNames]).value_counts().to_dict()
    print(f"  cases in file 02 by split: {counts}")
    if counts.get("unassigned"):
        print("  NOTE: 'unassigned' cases are in file 02 but in none of files 03-08; they are not used for training.")


def splitOf(name):
    i = caseNumber(name)
    if SPLIT_MAP:
        return SPLIT_MAP.get(i, "unassigned")
    if i in TRAIN_IDS:
        return "train"
    if i in VAL_IDS:
        return "validation"
    if i in TEST_IDS:
        return "test"
    return "unassigned"


def printDataChecks(data, params):
    """A short check of the data before it is trusted."""
    print("\n=== DATA CHECK ===")
    print(f"settings: {len(params)} cases x {list(params.columns)}")
    for key, df in data.items():
        cc = caseCols(df)
        empty = [c for c in cc if df[c].notna().sum() == 0]
        print(f"  {key:10s} cases={len(cc)}  history points={df[HISTORY_COL].notna().sum()}  "
              f"dates {df.index.min().date()} -> {df.index.max().date()}"
              + (f"  EMPTY cases: {empty[:5]}" if empty else ""))
    histEnd = data["oil_cum"][HISTORY_COL].dropna().index.max()
    if any(df.index.max() > histEnd + pd.Timedelta(days=1) for df in data.values()):
        print("  NOTE: some sheets hold dates AFTER the history end - check for future data.")
    ids = [c for c in caseCols(data["oil_cum"]) if c in params.index]
    print(f"  cases with both settings and curves: {len(ids)}  "
          f"splits: {pd.Series([splitOf(c) for c in ids]).value_counts().to_dict()}")


# =============================================================================
# STEP 2 - PREPARE CURVES: every case on the SAME dates as the real history
# =============================================================================
def toT(x, kind):
    """Optional transform before scoring/modelling ('log' keeps small water visible)."""
    x = np.asarray(x, float)
    return np.log10(np.clip(x, 0, None) + 1.0) if kind == "log" else x


def fromT(y, kind):
    y = np.asarray(y, float)
    return 10.0 ** y - 1.0 if kind == "log" else y


def evalCase(series, dates, isCum, t0):
    """Value of ONE case's curve at `dates`. Cumulative: linear between reported
    points, 0 at the start. Rate: nearest reported value. Outside the case's
    range the result is NaN (no extrapolation)."""
    s = series.dropna()
    if len(s) < 2:
        return np.full(len(dates), np.nan)

    def ns(idx):
        return pd.DatetimeIndex(idx).values.astype("datetime64[ns]").astype("int64").astype(float)

    x, y, xi = ns(s.index), s.values.astype(float), ns(dates)
    if isCum:
        if x[0] > t0:
            x, y = np.r_[t0, x], np.r_[0.0, y]
        return np.interp(xi, x, y, left=np.nan, right=np.nan)
    hi = np.clip(np.searchsorted(x, xi), 0, len(x) - 1)
    lo = np.clip(hi - 1, 0, len(x) - 1)
    pick = np.where(np.abs(xi - x[lo]) <= np.abs(x[hi] - xi), lo, hi)
    out = y[pick]
    out[(xi < x[0]) | (xi > x[-1])] = np.nan
    return out


def valueAt(series, dates, key, t0):
    return evalCase(series, pd.DatetimeIndex(dates), key.endswith("_cum"), t0)


def buildTargets(data, cases, t0):
    """For each scored stream: the real history and every case's curve on the SAME
    dates. Returns (targets, usable cases)."""
    T = {}
    for key, kind in MATCH_STREAMS.items():
        df = data[key]
        hist = df[HISTORY_COL].dropna()
        hist = hist[~hist.index.duplicated()]
        Y = np.array([toT(evalCase(df[c], hist.index, key.endswith("_cum"), t0), kind)
                      for c in cases])
        T[key] = dict(kind=kind, dates=hist.index, obs=toT(hist.values, kind), Y=Y)
    for t in T.values():                              # keep dates nearly every case covers
        keep = np.isfinite(t["Y"]).mean(axis=0) >= 0.95
        t["Y"], t["obs"], t["dates"] = t["Y"][:, keep], t["obs"][keep], t["dates"][keep]
    ok = np.ones(len(cases), bool)
    for t in T.values():
        ok &= np.isfinite(t["Y"]).all(axis=1)         # drop cases with holes
    for t in T.values():
        t["Y"] = t["Y"][ok]
        t["scale"] = max(float(np.abs(t["obs"]).max()), 1e-9)
    dropped = [c for c, o in zip(cases, ok) if not o]
    if dropped:
        print(f"   {len(dropped)} cases dropped (holes on the history dates): {dropped[:8]}")
    return T, [c for c, o in zip(cases, ok) if o]


def streamRmse(pred, t):
    """Error of predicted curves vs the real history, as a share of its largest value."""
    return np.sqrt(np.mean((pred - t["obs"][None, :]) ** 2, axis=1)) / t["scale"]


def totalMisfit(preds, T):
    """Sum over the scored streams (lower = closer to the real field)."""
    return sum(streamRmse(preds[k], T[k]) for k in T)


# =============================================================================
# STEP 3 - THE PROXY: PCA (compress each curve) + Gaussian process (settings -> curve)
# =============================================================================
class CurveProxy:
    """4 settings in -> the scored curves out. Per stream: PCA squeezes the curves into
    a few shape numbers, then one Gaussian process predicts each number. A Gaussian
    process suits ~70-100 runs: it needs few samples and reports its own uncertainty."""

    def __init__(self, paramCols, bounds):
        self.cols, self.bounds = list(paramCols), bounds

    def toUnit(self, P):
        """Physical settings -> [0,1] box (log scale for LOG_PARAMS)."""
        P = np.array(P, float, ndmin=2)
        U = np.empty_like(P)
        for j, c in enumerate(self.cols):
            lo, hi = self.bounds[c]
            if c in LOG_PARAMS:
                U[:, j] = (np.log(P[:, j]) - np.log(lo)) / (np.log(hi) - np.log(lo))
            else:
                U[:, j] = (P[:, j] - lo) / (hi - lo)
        return U

    def fromUnit(self, U):
        U = np.array(U, float, ndmin=2)
        P = np.empty_like(U)
        for j, c in enumerate(self.cols):
            lo, hi = self.bounds[c]
            if c in LOG_PARAMS:
                P[:, j] = np.exp(np.log(lo) + U[:, j] * (np.log(hi) - np.log(lo)))
            else:
                P[:, j] = lo + U[:, j] * (hi - lo)
        return P

    def fit(self, P, Y):
        U = self.toUnit(P)
        self.models = {}
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            for key, y in Y.items():
                mu = y.mean(axis=0)
                npc = int(min(PCA_MAX, y.shape[0] - 1, y.shape[1]))
                pca = PCA(n_components=npc).fit(y - mu)
                cum = np.nan_to_num(np.cumsum(pca.explained_variance_ratio_))
                k = int(min(npc, np.searchsorted(cum, PCA_VAR) + 1))
                Z = pca.transform(y - mu)[:, :k]
                gps = []
                for j in range(k):
                    kern = (ConstantKernel(1.0, (1e-2, 1e2))
                            * Matern(np.ones(U.shape[1]), (0.05, 50.0), nu=2.5)
                            + WhiteKernel(1e-4, (1e-8, 1e-1)))
                    gps.append(GaussianProcessRegressor(
                        kern, normalize_y=True, n_restarts_optimizer=2,
                        random_state=RANDOM_SEED).fit(U, Z[:, j]))
                self.models[key] = dict(mu=mu, comp=pca.components_[:k], gps=gps)
        return self

    def predict(self, P):
        U = self.toUnit(P)
        out = {}
        for key, m in self.models.items():
            Z = np.column_stack([gp.predict(U) for gp in m["gps"]])
            out[key] = m["mu"] + Z @ m["comp"]
        return out


def cvPredict(P, Y, paramCols, bounds, k=5):
    """Out-of-fold predictions: every case is predicted by a proxy that never saw it."""
    folds = np.arange(len(P)) % k
    np.random.RandomState(RANDOM_SEED).shuffle(folds)
    out = {key: np.zeros_like(y) for key, y in Y.items()}
    for f in range(k):
        tr, te = folds != f, folds == f
        pred = CurveProxy(paramCols, bounds).fit(P[tr], {key: y[tr] for key, y in Y.items()}).predict(P[te])
        for key in Y:
            out[key][te] = pred[key]
    return out


def verdictWord(errPct):
    """Plain-language grade for an error given as % of the largest history value."""
    return "excellent" if errPct < 2 else "good" if errPct < 5 else "fair" if errPct < 10 else "poor"


def proxyReport(pred, Y, T, label):
    """Accuracy of predicted curves vs the real simulated curves for these cases."""
    rows = []
    for key in T:
        err = np.sqrt(np.mean((pred[key] - Y[key]) ** 2, axis=1)) / T[key]["scale"]
        lt, lp = Y[key][:, -1], pred[key][:, -1]
        sst = np.sum((lt - lt.mean()) ** 2)
        rows.append({"dataset": label, "measure": key, "cases": len(err),
                     "typical_error_pct": 100 * err.mean(), "worst_error_pct": 100 * err.max(),
                     "R2_final": 1 - np.sum((lt - lp) ** 2) / sst if sst > 0 else np.nan})
    return pd.DataFrame(rows)


# =============================================================================
# STEP 4 - HISTORY MATCH: search ALL possible settings through the proxy
# =============================================================================
def runHistoryMatch(data, params):
    print("\n=== PROXY + HISTORY MATCH ===")
    pcols = list(params.columns)
    first = next(iter(MATCH_STREAMS))
    cases = [c for c in caseCols(data[first]) if c in params.index and params.loc[c].notna().all()]
    t0 = pd.Timestamp(SIM_START or min(data[k][HISTORY_COL].dropna().index.min()
                                       for k in MATCH_STREAMS)).value
    T, cases = buildTargets(data, cases, t0)
    P = params.loc[cases, pcols].values.astype(float)
    split = np.array([splitOf(c) for c in cases])
    bounds = {c: (params[c].min(), params[c].max()) for c in pcols}      # the tested ranges
    Yall = {k: t["Y"] for k, t in T.items()}
    print("cases usable: " + ", ".join(f"{s}={int((split == s).sum())}"
                                       for s in ["train", "validation", "test", "unassigned"]))
    tr = split == "train"
    if tr.sum() < 15:
        print("   WARNING: fewer than 15 train cases - all cases treated as train.")
        tr = np.ones(len(cases), bool)

    def sub(m):
        return {k: y[m] for k, y in Yall.items()}

    # ---- how accurate is the proxy? (cases it never saw) ----
    cvp = cvPredict(P[tr], sub(tr), pcols, bounds)
    reports = [proxyReport(cvp, sub(tr), T, "train (cross-validation)")]
    proxy = CurveProxy(pcols, bounds).fit(P[tr], sub(tr))
    held = {}
    for name in ["validation", "test"]:
        m = split == name
        if m.any():
            held[name] = (m, proxy.predict(P[m]))
            reports.append(proxyReport(held[name][1], sub(m), T, name))
    rep = pd.concat(reports, ignore_index=True)
    rep["verdict"] = rep["typical_error_pct"].map(verdictWord)
    print(rep.round(3).to_string(index=False))
    refM, refPred = held["validation"] if "validation" in held else (tr, cvp)
    refTrue = {k: y[refM] for k, y in Yall.items()}
    proxyErr = float(np.mean(np.abs(totalMisfit(refTrue, T) - totalMisfit(refPred, T))))
    print(f"proxy's own error in misfit units: {proxyErr:.4f}")
    use = tr | (split == "validation") if FINAL_REFIT_ON_VAL else tr
    if FINAL_REFIT_ON_VAL and (split == "validation").any():
        proxy = CurveProxy(pcols, bounds).fit(P[use], sub(use))
        print(f"final proxy refit on train+validation ({int(use.sum())} cases); test untouched")

    # ---- search every setting, not just the existing runs ----
    d = len(pcols)
    U = qmc.Sobol(d=d, scramble=True, seed=RANDOM_SEED).random_base2(N_SOBOL_POW2)
    samples = proxy.fromUnit(U)
    mis = totalMisfit(proxy.predict(samples), T)

    def objective(x):
        x = np.asarray(x, float)
        one = x.ndim == 1
        m = totalMisfit(proxy.predict(proxy.fromUnit(x.reshape(d, -1).T)), T)
        return float(m[0]) if one else m

    res = differential_evolution(objective, [(0.0, 1.0)] * d, seed=RANDOM_SEED, maxiter=150,
                                 popsize=20, tol=1e-8, vectorized=True,
                                 updating="deferred", polish=True)
    bestP = proxy.fromUnit(np.clip(res.x, 0, 1))[0]
    bestM = float(objective(np.clip(res.x, 0, 1)))
    if mis.min() < bestM:
        bestP, bestM = samples[int(mis.argmin())], float(mis.min())
    bestPred = proxy.predict(bestP[None, :])

    # ---- every setting that fits "about as well as the best" ----
    margin = max(ACCEPT_MARGIN_MIN, ACCEPT_MARGIN_REL * bestM, 2 * proxyErr)
    acc = mis <= bestM + margin
    A = samples[acc]
    existing = pd.Series(totalMisfit(Yall, T), index=cases).sort_values()
    print(f"\nBEST settings (misfit {bestM:.4f}); acceptable = within {margin:.4f} of the best: "
          f"{int(acc.sum())} of {len(mis)} trial sets")
    print(f"best EXISTING simulated case: {existing.index[0]} (misfit {existing.iloc[0]:.4f})")

    Ua, rows = proxy.toUnit(A), []
    for j, c in enumerate(pcols):
        w = float(Ua[:, j].max() - Ua[:, j].min())
        verdict = "PINNED DOWN" if w < 0.5 else "partly constrained" if w < 0.8 else "NOT constrained"
        rows.append({"setting": c, "best": bestP[j], "acceptable_min": A[:, j].min(),
                     "acceptable_max": A[:, j].max(), "tested_min": bounds[c][0],
                     "tested_max": bounds[c][1], "share_of_range": w, "verdict": verdict})
        print(f"  {verdict:18s} {c}: best {bestP[j]:.4g}, acceptable {A[:, j].min():.3g} to "
              f"{A[:, j].max():.3g} (tested {bounds[c][0]:.3g} to {bounds[c][1]:.3g})")
    ident = pd.DataFrame(rows)

    resid = {}
    for k, t in T.items():
        r = (bestPred[k][0] - t["obs"]) / t["scale"]
        resid[k] = float(np.sqrt(np.mean(r ** 2)))
    print("misfit at the best settings (RMSE as % of the largest history value): "
          + ", ".join(f"{k} {100 * v:.1f}%" for k, v in resid.items()))
    if max(resid.values()) > 0.10:
        print("   CHECK: the proxy cannot reproduce history in at least one stream; the 4 settings "
              "may not span the truth.")
    return dict(proxy=proxy, bestParams=bestP, bestMisfit=bestM, acceptable=A, accMis=mis[acc],
                targets=T, Yall=Yall, held=held, rep=rep, ident=ident, bestPred=bestPred,
                resid=resid, existing=existing, t0=t0, pcols=pcols, bounds=bounds,
                samples=samples, mis=mis, cases=cases, proxyErr=proxyErr)



def rankCases(params, hm):
    """All existing simulated cases ranked by closeness to the real history (one score:
    the history-match misfit). Also keeps each stream's error so the ranking can be explained."""
    T = hm["targets"]
    tab = pd.DataFrame({"misfit": hm["existing"]})
    tab.insert(0, "rank", np.arange(1, len(tab) + 1))
    for k in T:
        tab[f"err_{k}"] = pd.Series(100 * streamRmse(hm["Yall"][k], T[k]), index=hm["cases"])
    top5 = ", ".join(tab.index[:5])
    print(f"\nCLOSEST EXISTING SIMULATED CASES: {top5}")
    return tab.join(params[hm["pcols"]])


# =============================================================================
# STEP 5 - FORECAST 2008 -> 2028 (EXTRAPOLATION beyond the data)
# =============================================================================
def arps(t, qi, Di, b):
    """Hyperbolic decline: the rate falls fast at first, then more slowly."""
    return qi / np.power(1.0 + b * Di * t, 1.0 / b)


def fitArps(t, q):
    """Fit (qi, Di, b) to a rate history; t = years from the start of the window."""
    m = np.isfinite(t) & np.isfinite(q) & (q > 0)
    if m.sum() < 5:
        return None
    t, q = t[m], q[m]
    lq = np.log(q)

    def f(tt, lqi, lDi, b):
        return np.log(arps(tt, np.exp(lqi), np.exp(lDi), b))

    best = None
    for b0 in (0.1, 0.5, 0.9):
        try:
            popt, _ = curve_fit(f, t, lq, p0=[lq[0], np.log(0.15), b0],
                                bounds=([lq.min() - 3, -9, 0.01], [lq.max() + 3, 2.0, 1.0]),
                                maxfev=4000)
            sse = np.sum((f(t, *popt) - lq) ** 2)
            if best is None or sse < best[0]:
                best = (sse, popt)
        except Exception:
            continue
    if best is None:                                    # fallback: plain exponential
        slope, icpt = np.polyfit(t, lq, 1)
        return (float(np.exp(icpt)), float(max(-slope, 1e-6)), 0.01)
    return (float(np.exp(best[1][0])), float(np.exp(best[1][1])), float(best[1][2]))


def fitShapes(data, case, cut, windowYears, t0):
    """Everything the forecast needs from one case using history up to `cut`:
    its state at `cut` and decline shapes fitted on the last `windowYears`."""
    w0 = cut - pd.Timedelta(days=int(365.25 * windowYears))
    out = {"w0": w0, "cut": cut}
    for s in ("oil", "gas"):
        ser = data[f"{s}_rate"][case].dropna()
        ser = ser[(ser.index >= w0) & (ser.index <= cut)]
        out[s] = fitArps((ser.index - w0).days.values / 365.25, ser.values)
        if out[s] is None:
            return None
    st = {k: float(valueAt(data[k][case], [cut], k, t0)[0])
          for k in ["oil_cum", "gas_cum", "water_cum", "oil_rate", "gas_rate", "water_rate"]}
    if not np.all(np.isfinite(list(st.values()))):
        return None
    out["state"] = st
    grid = pd.date_range(w0, cut, periods=24)
    qo = valueAt(data["oil_rate"][case], grid, "oil_rate", t0)
    qw = valueAt(data["water_rate"][case], grid, "water_rate", t0)
    npv = valueAt(data["oil_cum"][case], grid, "oil_cum", t0)
    with np.errstate(divide="ignore", invalid="ignore"):
        wor = np.where(qo > 0, qw / qo, np.nan)
    ok = np.isfinite(wor) & np.isfinite(npv) & (wor > 1e-4)
    slopeW = 0.0
    if ok.sum() >= 5:
        slope, _ = np.polyfit(npv[ok] / max(st["oil_cum"], 1.0), np.log(wor[ok]), 1)
        slopeW = max(float(slope), 0.0) if np.isfinite(slope) else 0.0   # water cut does not fall with production
    out["worSlope"] = slopeW
    return out


def project(shp, outDates):
    """Cumulative gas / oil / water at `outDates` (all on/after shp['cut'])."""
    cut, st = shp["cut"], shp["state"]
    outDates = pd.DatetimeIndex(outDates)
    d = np.arange(0, (outDates.max() - cut).days + 8, 7.0)           # weekly steps
    off = (cut - shp["w0"]).days / 365.25
    ty, dd = off + d / 365.25, np.diff(d)

    def integrate(q):
        return np.concatenate([[0.0], np.cumsum(0.5 * (q[1:] + q[:-1]) * dd)])

    q = {}
    for s in ("oil", "gas"):                            # rescaled to join the last real rate
        qi, Di, b = shp[s]
        q[s] = arps(ty, qi, Di, b) * (st[f"{s}_rate"] / max(arps(off, qi, Di, b), 1e-12))
    cumO = st["oil_cum"] + integrate(q["oil"])
    worEnd = max(st["water_rate"] / max(st["oil_rate"], 1e-9), 1e-4)
    wor = np.minimum(worEnd * np.exp(np.minimum(shp["worSlope"] * (cumO - st["oil_cum"])
                                                / max(st["oil_cum"], 1.0), 50.0)), max(WOR_MAX, worEnd))
    x = (outDates - cut).days.values.astype(float)
    return {"gas": np.interp(x, d, st["gas_cum"] + integrate(q["gas"])),
            "oil": np.interp(x, d, cumO),
            "water": np.interp(x, d, st["water_cum"] + integrate(wor * q["oil"]))}


def hindcast(data, cases, windowYears, histEnd, t0):
    """Hide the last HINDCAST_YEARS of history, forecast them, and compare the predicted
    production INCREASE with what the simulator really did."""
    cut = histEnd - pd.Timedelta(days=int(365.25 * HINDCAST_YEARS))
    rows = {}
    for c in cases:
        shp = fitShapes(data, c, cut, windowYears, t0)
        if shp is None:
            continue
        pred, r = project(shp, [histEnd]), {}
        for s in STREAMS:
            trueEnd = float(valueAt(data[f"{s}_cum"][c], [histEnd], f"{s}_cum", t0)[0])
            inc = trueEnd - shp["state"][f"{s}_cum"]
            r[s] = (pred[s][0] - shp["state"][f"{s}_cum"]) / inc - 1 if inc > 0 else np.nan
        rows[c] = r
    return pd.DataFrame(rows).T


def selectCases(U, oilEnd, n, mode):
    """Indices into the acceptable settings (sorted best-misfit first)."""
    n = min(n, len(U))
    if mode == "best":
        return list(range(n))
    if mode == "span_forecast":
        order = np.argsort(oilEnd)
        pick = np.unique(np.round(np.linspace(0.03, 0.97, n) * (len(order) - 1)).astype(int))
        return sorted(order[pick].tolist())
    chosen = [0]                                        # best match, then the most different
    while len(chosen) < n:
        dist = np.min(np.linalg.norm(U[:, None, :] - U[chosen][None, :, :], axis=2), axis=1)
        dist[chosen] = -1
        chosen.append(int(np.argmax(dist)))
    return chosen


# ---- the submission template ------------------------------------------------
def findTemplate():
    path = findFile(FILES["template"])
    if path is None:
        for p in listWorkbooks():
            if "template" in p.stem.lower():
                return p
    return path


DATE_TEXT = re.compile(r"^\s*\d{1,4}\s*[/\-.]\s*\d{1,2}\s*[/\-.]\s*\d{1,4}\s*$|^\s*\d{1,2}[\s\-]+[A-Za-z]{3,9}[\s\-,]+\d{2,4}\s*$")


def cellDateKind(v):
    """'dt' = real date, 'txt' = text that looks like a date, 'num' = Excel serial number, None = other."""
    if isinstance(v, (datetime.datetime, datetime.date, pd.Timestamp)):
        return "dt"
    if isinstance(v, str) and DATE_TEXT.match(v):
        return "txt"
    if isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool) and 20000 < float(v) < 80000:
        return "num"
    return None


def _monotonic(ts):
    return len(ts) >= 2 and all(pd.notna(t) for t in ts) and all(b > a for a, b in zip(ts, ts[1:]))


_NOTED = set()


def resolveTemplateDates(cells):
    """cells = raw values of the template's date column. Returns Timestamps (same order).
    Handles real dates, text ('04/01/2008') and Excel serial numbers. Month-first is the default;
    day-first is used only if month-first does not give dates that increase down the column."""
    def one(v, dayfirst, swap):
        k = cellDateKind(v)
        if k == "dt":
            t = pd.Timestamp(v)
            t = t.tz_localize(None) if t.tzinfo else t
            if swap and t.day <= 12:                                # a day/month mix-up made by Excel
                t = pd.Timestamp(year=t.year, month=t.day, day=t.month)
            return t.normalize()
        if k == "num":
            return (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(v))).normalize()
        if k == "txt":
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                return pd.to_datetime(v, dayfirst=dayfirst, errors="coerce").normalize()
        return pd.NaT

    def day1(ts):                                         # template dates are normally the 1st of a month
        return np.mean([t.day == 1 for t in ts if pd.notna(t)]) if any(pd.notna(t) for t in ts) else 0.0

    options = [([one(v, False, False) for v in cells], "as written"),
               ([one(v, True, False) for v in cells], "day/month order (text dates)"),
               ([one(v, False, True) for v in cells], "day/month swapped")]
    for ts, label in options:                             # best: increasing AND on the 1st of the month
        if _monotonic(ts) and day1(ts) >= 0.9:
            if label != "as written" and label not in _NOTED:
                _NOTED.add(label)
                print(f"  NOTE: template dates read as {label}: that is the only reading that gives dates on the 1st "
                      "of each month.")
            return ts, label
    for ts, label in options:                             # otherwise: at least increasing
        if _monotonic(ts):
            return ts, label
    return options[0][0], "unchecked"


def fallbackSchedule():
    return pd.date_range(FORECAST_START_IF_NO_TEMPLATE, FORECAST_END_IF_NO_TEMPLATE, freq="QS")


def readTemplateLayout(path, allowFallback=True):
    """Find, in every 'Case N' sheet: header row, date column, the gas / oil / water
    cumulative columns and the (row, date) of each data row. Nothing is hard-coded.
    Dates may be real Excel dates, text, or serial numbers. If a sheet's dates cannot be
    read at all, the quarterly schedule from the CONFIG is used and written into the sheet."""
    wb = load_workbook(path)
    layout = {}
    for name in wb.sheetnames:
        if not re.search(r"case\s*_?\s*\d+", name, re.I):
            continue
        ws, hdr = wb[name], None
        for r in range(1, min(ws.max_row, 15) + 1):
            cells = [str(c.value).lower() if c.value is not None else "" for c in ws[r]]
            if sum("cumulative" in x for x in cells) >= 2:
                hdr = r
                break
        if hdr is None:
            continue
        cols = {}
        for c in ws[hdr]:
            v = str(c.value).lower() if c.value is not None else ""
            if "cumulative" in v:
                for s in STREAMS:
                    if s in v:
                        cols[s] = c.column
                        break
        cand = [col for col in range(1, ws.max_column + 1) if col not in cols.values()]
        counts = {col: sum(cellDateKind(ws.cell(row=r, column=col).value) is not None
                           for r in range(hdr + 1, ws.max_row + 1)) for col in cand}
        dateCol = max(counts, key=counts.get) if counts and max(counts.values()) > 0 else 1
        rowsRaw = [(r, ws.cell(row=r, column=dateCol).value) for r in range(hdr + 1, ws.max_row + 1)
                   if cellDateKind(ws.cell(row=r, column=dateCol).value) is not None]
        synthetic = False
        if len(rowsRaw) >= 2:
            ts, _ = resolveTemplateDates([v for _, v in rowsRaw])
            rows = [(r, t) for (r, _), t in zip(rowsRaw, ts) if pd.notna(t)]
        else:
            rows = []
        if len(rows) < 2:
            if not allowFallback:
                continue
            print(f"  WARNING: could not read the dates in template sheet '{name}'. Using the quarterly schedule "
                  f"{FORECAST_START_IF_NO_TEMPLATE} -> {FORECAST_END_IF_NO_TEMPLATE} and writing the dates into the sheet.")
            rows = [(hdr + 1 + i, d) for i, d in enumerate(fallbackSchedule())]
            synthetic = True
        layout[name] = {"hdr": hdr, "cols": cols, "dateCol": dateCol, "rows": rows, "synthetic": synthetic}
    return wb, dict(sorted(layout.items(), key=lambda kv: caseNumber(kv[0])))


def makeOwnTemplate(dates):
    """Fallback when no template file exists: same layout, our own dates."""
    path = OUTPUT_DIR / "own_template.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    for k in range(1, N_SELECT + 1):
        ws = wb.create_sheet(f"Case {k}")
        ws.append(["Date", "Gas production cumulative [MSCF]",
                   "Oil production cumulative [STB]", "Water production cumulative [STB]"])
        for d in dates:
            ws.append([d.to_pydatetime()])
    wb.save(path)
    return path


def checkDeliverable(path, obsEnd, histEnd=None):
    """Read the FILLED template back from disk and verify the deliverable's rules.
    Returns a table (one row per sheet) so it can also go into the summary Excel."""
    wb, layout = readTemplateLayout(path, allowFallback=False)
    print("\n=== DELIVERABLE CHECK (reading the saved file back) ===")
    rows = []
    for name, lay in layout.items():
        ws, problems = wb[name], []
        for s in STREAMS:
            vals = np.array([ws.cell(row=r, column=lay["cols"][s]).value
                             if ws.cell(row=r, column=lay["cols"][s]).value is not None else np.nan
                             for r, _ in lay["rows"]], float)
            if np.isnan(vals).any():
                problems.append(f"{s}: blank cells")
                continue
            if (vals < 0).any():
                problems.append(f"{s}: negative values")
            if (np.diff(vals) < -1e-6 * max(abs(vals).max(), 1)).any():
                problems.append(f"{s}: cumulative decreases")
            after = np.array([histEnd is None or d >= histEnd for _, d in lay["rows"]])
            if after.any() and np.isfinite(obsEnd.get(s, np.nan)) and vals[after][0] < obsEnd[s] * (1 - 1e-9):
                problems.append(f"{s}: starts below the 2008 history value")
        status = "PASS" if not problems else "FAIL: " + "; ".join(problems)
        print(f"  {name:8s} rows={len(lay['rows'])}  {status}")
        rows.append({"sheet": name, "rows": len(lay["rows"]), "result": status})
    return pd.DataFrame(rows)


def runForecast(data, params, hm):
    print("\n=== FORECAST 2008 -> 2028 ===")
    proxy, t0, pcols, bounds = hm["proxy"], hm["t0"], hm["pcols"], hm["bounds"]
    histEnd = data["oil_cum"][HISTORY_COL].dropna().index.max()
    cases = [c for c in caseCols(data["oil_cum"]) if c in params.index and params.loc[c].notna().all()]
    print(f"history ends {histEnd.date()}; the simulator curves stop here too, so everything "
          "after is extrapolation")

    # ---- the template defines the dates to forecast ----
    tpath = findTemplate()
    layout = None
    if tpath is not None:
        wb, layout = readTemplateLayout(tpath)
        if not layout:
            print(f"  WARNING: no 'Case N' sheet with 'cumulative' columns found in {tpath.name} -> writing our own "
                  "10-sheet workbook instead.")
            tpath = None
    if tpath is not None:
        fdates = pd.DatetimeIndex([d for _, d in next(iter(layout.values()))["rows"]])
        print(f"template: {tpath.name} ({len(layout)} case sheets), {len(fdates)} dates "
              f"{fdates.min().date()} -> {fdates.max().date()}   first dates: "
              + ", ".join(str(d.date()) for d in fdates[:4]))
        if (fdates <= histEnd).any():
            print(f"  NOTE: {int((fdates <= histEnd).sum())} template dates are on/before the history end "
                  f"({histEnd.date()}): those rows are filled with the real history instead of a forecast.")
            fdates = fdates[fdates > histEnd]
    else:
        print("no template file found -> writing our own 10-sheet workbook")
        fdates = pd.date_range(FORECAST_START_IF_NO_TEMPLATE, FORECAST_END_IF_NO_TEMPLATE, freq="QS")

    # ---- hindcast: test the method on the past, choose the fitting window ----
    errs = {w: hindcast(data, cases, w, histEnd, t0) for w in FIT_WINDOW_CANDIDATES}
    for w, e in errs.items():
        print(f"  hindcast, decline curves fitted on last {w:g} yr -> typical error of the "
              f"{HINDCAST_YEARS:g}-yr increase: "
              + "  ".join(f"{s} {e[s].abs().median():.1%}" for s in STREAMS))
    chosen = FIT_WINDOW_YEARS or min(
        errs, key=lambda w: np.nanmean([errs[w][s].abs().median() for s in STREAMS]))
    if chosen not in errs:
        errs[chosen] = hindcast(data, cases, chosen, histEnd, t0)
    print(f"  chosen fitting window: last {chosen:g} years")

    # ---- one extrapolated forecast per case, then a proxy that forecasts for ANY settings ----
    fdExt = pd.DatetimeIndex([histEnd]).append(fdates)        # column 0 = state at history end
    F, okCases = {s: [] for s in STREAMS}, []
    for c in cases:
        shp = fitShapes(data, c, histEnd, chosen, t0)
        if shp is None:
            continue
        pr = project(shp, fdExt)
        if not all(np.all(np.isfinite(pr[s])) for s in STREAMS):
            continue                                     # never let a bad case poison the proxy
        okCases.append(c)
        for s in STREAMS:
            F[s].append(pr[s])
    print(f"  cases with a usable forecast fit: {len(okCases)} of {len(cases)}")
    if len(okCases) < 10:
        raise RuntimeError("Fewer than 10 cases have a usable forecast. Usual cause: the simulated case curves do not "
                           f"reach the history end ({histEnd.date()}) or have holes there - check the case dates.")
    if len(okCases) < 0.8 * len(cases):
        print("  WARNING: many cases have no value at the history end date and were skipped - check the case dates.")
    P = params.loc[okCases, pcols].values.astype(float)
    split = np.array([splitOf(c) for c in okCases])
    use = np.ones(len(okCases), bool) if FORECAST_REFIT_ON_ALL else np.isin(split, ["train", "validation"])
    if use.sum() < 15:
        print("  WARNING: fewer than 15 train/validation cases - using ALL cases for the forecast proxy.")
        use = np.ones(len(okCases), bool)
    kinds = {f"fc_{s}": ("log" if s == "water" else "linear") for s in STREAMS}
    Yf = {f"fc_{s}": toT(np.array(F[s]), kinds[f"fc_{s}"]) for s in STREAMS}
    fproxy = CurveProxy(pcols, bounds).fit(P[use], {k: y[use] for k, y in Yf.items()})

    # ---- forecast for every ACCEPTABLE setting -> spread between equally good matches ----
    A, am = hm["acceptable"], hm["accMis"]
    if len(A) < max(MIN_FORECAST_SETS, N_SELECT):
        k = max(MIN_FORECAST_SETS, N_SELECT * 5)
        idx = np.argsort(hm["mis"])[:k]
        print(f"  only {len(A)} acceptable sets; using the {k} best-fitting trial sets instead")
        A, am = hm["samples"][idx], hm["mis"][idx]
    if len(A) > MAX_FORECAST_SETS:
        idx = np.random.RandomState(RANDOM_SEED).choice(len(A), MAX_FORECAST_SETS, replace=False)
        A, am = A[idx], am[idx]
    order = np.argsort(am)
    A, am = A[order], am[order]
    pred = fproxy.predict(A)
    obsEnd = {s: lastValid(data[f"{s}_cum"][HISTORY_COL]) for s in STREAMS}
    ens = {}
    for s in STREAMS:
        arr = fromT(pred[f"fc_{s}"], kinds[f"fc_{s}"])
        fc = arr[:, 1:]
        if ANCHOR_TO_HISTORY and np.isfinite(obsEnd[s]):
            fc = fc + (obsEnd[s] - arr[:, 0])[:, None]          # join the real history exactly
        fc = np.maximum.accumulate(np.maximum(fc, 0), axis=1)
        if np.isfinite(obsEnd[s]):
            fc = np.maximum(fc, obsEnd[s])
        ens[s] = fc
    okRows = np.all([np.all(np.isfinite(ens[s]), axis=1) for s in STREAMS], axis=0)
    if not okRows.any():
        raise RuntimeError("The forecast came out empty / not-a-number for every setting. Check the production "
                           "sheets for blanks around the history end date.")
    if not okRows.all():
        print(f"  WARNING: {int((~okRows).sum())} of {len(okRows)} forecast sets were not finite and were dropped.")
        A, am = A[okRows], am[okRows]
        ens = {s: ens[s][okRows] for s in STREAMS}
    bands = {s: tuple(np.percentile(ens[s], q, axis=0) for q in (10, 50, 90)) for s in STREAMS}
    print("  forecast check (middle, at the last date): " + "   ".join(
        f"{s} {bands[s][1][-1]:,.0f} {UNITS[s]}" for s in STREAMS))

    # ---- choose the cases to submit ----
    pick = selectCases(proxy.toUnit(A), ens["oil"][:, -1], N_SELECT, SELECTION_MODE)
    pick = sorted(pick, key=lambda i: am[i])
    rows = []
    for rank, i in enumerate(pick, start=1):
        row = {"case": rank, "misfit": am[i], **dict(zip(pcols, A[i]))}
        for s in STREAMS:
            row[f"{s}_cum_{fdates[-1].year}"] = ens[s][i, -1]
        rows.append(row)
    sel = pd.DataFrame(rows)
    print(f"\nSELECTED CASES (mode '{SELECTION_MODE}', {len(pick)} of {len(A)} acceptable sets):")
    print(sel.round(4).to_string(index=False))

    # ---- fill a COPY of the template and check it ----
    if tpath is None:
        tpath = makeOwnTemplate(fdates)
        wb, layout = readTemplateLayout(tpath)
    outPath = OUTPUT_DIR / f"{tpath.stem}_FILLED.xlsx"
    sheets = list(layout.items())
    if len(sheets) < len(pick):
        print(f"WARNING: template has {len(sheets)} case sheets but {len(pick)} cases were selected.")
    # Every template date gets a value by position on the forecast curve; dates on/before the
    # history end get the real history. Nothing depends on dates matching exactly.
    ns = lambda idx: pd.DatetimeIndex(idx).values.astype("datetime64[ns]").astype("int64").astype(float)
    xF = ns(pd.DatetimeIndex([histEnd]).append(fdates))
    histCum = {s: data[f"{s}_cum"][HISTORY_COL].dropna() for s in STREAMS}
    written = 0
    for (name, lay), i in zip(sheets, pick):
        ws = wb[name]
        order = sorted(lay["rows"], key=lambda rd: rd[1])
        for s in STREAMS:
            vals = []
            for r, d in order:
                if d <= histEnd:                                    # real history, never above the last real value
                    v = min(float(np.interp(ns([d])[0], ns(histCum[s].index), histCum[s].values)), obsEnd[s])
                else:
                    v = float(np.interp(ns([d])[0], xF, np.concatenate([[obsEnd[s]], ens[s][i]])))
                vals.append(v)
            vals = np.maximum.accumulate(np.array(vals))            # cumulative can never decrease
            if not np.all(np.isfinite(vals)):
                raise RuntimeError(f"non-finite values for {name}, {s}")
            for (r, d), v in zip(order, vals):
                ws.cell(row=r, column=lay["cols"][s]).value = float(v)
                written += 1
        if lay.get("synthetic"):
            for r, d in order:
                ws.cell(row=r, column=lay["dateCol"], value=d.to_pydatetime()).number_format = "mm/dd/yyyy"
    if written == 0:
        raise RuntimeError("Nothing was written to the template: no date rows were found in its case sheets.")
    wb.save(outPath)
    print(f"\nsaved filled template: {outPath}   ({written} cells over {min(len(sheets), len(pick))} sheets; "
          "original template untouched)")
    check = checkDeliverable(outPath, obsEnd, histEnd)
    return dict(fdates=fdates, bands=bands, histEnd=histEnd, obsEnd=obsEnd, errs=errs, chosen=chosen,
                selected=sel, selCurves={s: [ens[s][i] for i in pick] for s in STREAMS},
                check=check, templatePath=outPath, nAcceptable=len(A))


# =============================================================================
# STEP 6 - THE GRAPHS (each headline states its conclusion; caption = one line)
# =============================================================================
SHORT = {"oil_cum": "Oil cum. error [%]", "water_cum": "Water cum. error [%]", "water_rate": "Water rate error [%]"}
SPLIT_COL = {"train": BLUE, "validation": "#ff7f0e", "test": "#9467bd"}


def savePlot(fig, name):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    fig.savefig(OUTPUT_DIR / name, dpi=140, bbox_inches="tight")
    print(f"saved {OUTPUT_DIR / name}")


def millions(ax):
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v / 1e6:.0f}M"))


def caption(fig, text):
    fig.supxlabel(text, fontsize=11, color="#444444")


def plotProblem(hm):
    """1: the problem - 100 simulated reservoirs, one real field."""
    T = hm["targets"]
    fig, axes = plt.subplots(1, len(T), figsize=(6 * len(T), 5), constrained_layout=True)
    for ax, (key, t) in zip(np.atleast_1d(axes), T.items()):
        for y in t["Y"]:
            ax.plot(t["dates"], fromT(y, t["kind"]), color="#bbbbbb", lw=0.7)
        ax.plot(t["dates"], fromT(t["obs"], t["kind"]), "ko", ms=4, label="real field")
        ax.plot([], [], color="#bbbbbb", label=f"{len(t['Y'])} simulated cases")
        ax.set_title(PLAIN[key], weight="bold")
        ax.set_ylabel(UNITS[key])
        ax.tick_params(axis="x", rotation=30)
        ax.grid(alpha=0.3)
        if key.endswith("_cum"):
            millions(ax)
    np.atleast_1d(axes)[0].legend(loc="upper left")
    fig.suptitle("The problem: many possible reservoirs (grey), one real field (black dots)",
                 fontsize=15, weight="bold")
    caption(fig, "Each grey line is one simulation with a different mix of 4 uncertain settings. "
                 "History matching = find the settings whose curve lands on the dots.")
    savePlot(fig, "1_the_problem.png")


def plotProxyTrust(hm):
    """2: can we trust the fast ML stand-in (proxy)?"""
    T, held, Yall, rep = hm["targets"], hm["held"], hm["Yall"], hm["rep"]
    heldRep = rep[rep["dataset"].isin(["validation", "test"])]
    worst = heldRep.groupby("measure")["typical_error_pct"].mean().max() if len(heldRep) else np.nan
    head = ("Yes - the proxy predicts unseen cases accurately" if worst < 2 else
            "Mostly - the proxy predicts unseen cases well, with some error" if worst < 5 else
            "With caution - the proxy has noticeable error on unseen cases")
    fig, axes = plt.subplots(1, len(T), figsize=(6 * len(T), 5.4), constrained_layout=True)
    for ax, (key, t) in zip(np.atleast_1d(axes), T.items()):
        kind, vals = t["kind"], []
        for name in ["validation", "test"]:
            if name in held:
                m, pred = held[name]
                tru, prd = fromT(Yall[key][m][:, -1], kind), fromT(pred[key][:, -1], kind)
                ax.scatter(tru, prd, s=50, color=SPLIT_COL[name], edgecolor="k", label=f"{name} cases", zorder=3)
                vals += list(tru) + list(prd)
        if vals:
            ax.plot([min(vals), max(vals)], [min(vals), max(vals)], "k--", lw=1.2, label="perfect prediction")
        r = heldRep[heldRep["measure"] == key]
        err = float(r["typical_error_pct"].mean()) if len(r) else np.nan
        word = verdictWord(err) if np.isfinite(err) else "n/a"
        ax.set_title(f"{PLAIN[key]}\ntypical error {err:.1f}% ({word})", weight="bold",
                     color={"excellent": GREEN, "good": GREEN, "fair": AMBER, "poor": RED}.get(word, "k"))
        ax.set_xlabel(f"real simulator value at end of history [{UNITS[key]}]")
        ax.set_ylabel("proxy prediction")
        ax.grid(alpha=0.3)
        if key.endswith("_cum"):
            millions(ax)
            ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v / 1e6:.0f}M"))
    np.atleast_1d(axes)[0].legend(fontsize=9)
    fig.suptitle(f"Can we trust the proxy?  {head}", fontsize=15, weight="bold")
    caption(fig, "Each dot is a case the proxy never saw in training. On the dashed line = proxy equals the "
                 f"real simulator. Error = % of the largest history value.\nTrain / validation / test: {SPLIT_SOURCE}.")
    savePlot(fig, "2_can_we_trust_the_proxy.png")


def plotRanking(hm, rank):
    """3: which existing simulated case is closest to the real history, and can the proxy beat it?"""
    top = rank.head(10)
    names = ["Best settings found\n(proxy, not a run)"] + list(top.index)
    vals = [hm["bestMisfit"]] + list(top["misfit"])
    cols = ["#d62728"] + [BLUE] * len(top)
    y = np.arange(len(vals))[::-1]
    fig, ax = plt.subplots(figsize=(11, 6), constrained_layout=True)
    ax.barh(y, vals, color=cols, edgecolor="k", height=0.7)
    for yy, v in zip(y, vals):
        ax.text(v, yy, f"  {v:.3f}", va="center", fontsize=10)
    ax.set_yticks(y)
    ax.set_yticklabels(names)
    ax.set_xlabel("misfit = oil + water cumulative + water-rate error (lower = closer to the real field)")
    ax.set_xlim(0, max(vals) * 1.15)
    ax.grid(alpha=0.3, axis="x")
    ax.legend(handles=[Patch(fc=c, ec="k", label=l) for c, l in
                       [("#d62728", "proxy prediction (not a simulator run)"),
                        (BLUE, "existing simulated case")]], loc="upper right", fontsize=9)
    better = hm["bestMisfit"] < vals[1]
    fig.suptitle(f"Closest of the {len(rank)} simulated cases: {top.index[0]} (misfit {vals[1]:.3f}).  "
                 + (f"The proxy search finds settings that fit better ({hm['bestMisfit']:.3f})." if better
                    else "The proxy search does not beat it."), fontsize=14, weight="bold")
    caption(fig, "Top 10 existing cases by misfit. The red bar is the proxy's prediction for the best settings "
                 "found, so it has not been run in the real simulator.")
    savePlot(fig, "3_closest_cases.png")


def plotWhatHistoryTells(hm):
    """4: which settings does the real history pin down?"""
    ident, proxy = hm["ident"], hm["proxy"]
    pinned = list(ident.loc[ident["verdict"] == "PINNED DOWN", "setting"])
    free = list(ident.loc[ident["verdict"] == "NOT constrained", "setting"])
    head = ("History pins down: " + (", ".join(pinned) if pinned else "none of the settings")
            + (".  It cannot tell: " + ", ".join(free) if free else "."))
    col = {"PINNED DOWN": GREEN, "partly constrained": AMBER, "NOT constrained": RED}

    def position(setting, value):                         # 0..1 inside the tested range
        row = np.array(hm["bestParams"], float)
        j = hm["pcols"].index(setting)
        row[j] = value
        return float(proxy.toUnit(row[None, :])[0, j])

    fig, ax = plt.subplots(figsize=(13, 1.0 * len(ident) + 1.8), constrained_layout=True)
    for i, r in enumerate(ident.iloc[::-1].itertuples()):
        lo, hi, b = (position(r.setting, v) for v in (r.acceptable_min, r.acceptable_max, r.best))
        ax.barh(i, 1, color="#e6e6e6", edgecolor="k", height=0.5)
        ax.barh(i, hi - lo, left=lo, color=col[r.verdict], edgecolor="k", height=0.5)
        ax.plot(b, i, "k*", ms=16, zorder=5)
        ax.text(0.5, i + 0.36, f"best {r.best:.3g}   |   plausible {r.acceptable_min:.3g} to "
                f"{r.acceptable_max:.3g}   |   tested {r.tested_min:.3g} to {r.tested_max:.3g}",
                ha="center", va="bottom", fontsize=9, color="#333333")
    ax.set_yticks(range(len(ident)))
    ax.set_yticklabels([f"{r.setting}\n{r.verdict}" for r in ident.iloc[::-1].itertuples()], fontsize=10)
    for lab, r in zip(ax.get_yticklabels(), ident.iloc[::-1].itertuples()):
        lab.set_color(col[r.verdict])
    ax.set_xlim(0, 1)
    ax.set_ylim(-0.5, len(ident) - 0.1)
    ax.set_xticks([0, 1])
    ax.set_xticklabels(["lowest value tested", "highest value tested"])
    fig.suptitle(head, fontsize=14, weight="bold")
    caption(fig, "Grey = every value tested. Colour = values that still match history about as well as the best. "
                 "Narrow = history determines it. Star = best value.")
    savePlot(fig, "4_what_history_tells_us.png")


def plotHistoryMatch(hm):
    """5: the best settings against the real field."""
    T, proxy, A, bestPred, resid = hm["targets"], hm["proxy"], hm["acceptable"], hm["bestPred"], hm["resid"]
    rng = np.random.RandomState(RANDOM_SEED)
    ens = proxy.predict(A[rng.choice(len(A), min(len(A), 300), replace=False)])
    avg = 100 * np.mean(list(resid.values()))
    topCase = hm["existing"].index[0]
    fig, axes = plt.subplots(1, len(T), figsize=(6 * len(T), 5), constrained_layout=True)
    for ax, (key, t) in zip(np.atleast_1d(axes), T.items()):
        kind, x = t["kind"], t["dates"]
        lo, hi = np.percentile(ens[key], [10, 90], axis=0)
        ax.fill_between(x, fromT(lo, kind), fromT(hi, kind), color="#d62728", alpha=0.25,
                        label="other settings that fit about as well")
        ax.plot(x, fromT(bestPred[key][0], kind), color="#d62728", lw=2.5, label="best settings found (proxy)")
        ax.plot(x, fromT(t["Y"][hm["cases"].index(topCase)], kind), color="#2ca02c", lw=1.8, ls="--",
                label=f"closest existing case ({topCase})")
        ax.plot(x, fromT(t["obs"], kind), "ko", ms=4, label="real field")
        ax.set_title(f"{PLAIN[key]}\nmismatch {100 * resid[key]:.1f}% of peak", weight="bold")
        ax.set_ylabel(UNITS[key])
        ax.tick_params(axis="x", rotation=30)
        ax.grid(alpha=0.3)
        if key.endswith("_cum"):
            millions(ax)
    np.atleast_1d(axes)[0].legend(fontsize=9, loc="upper left")
    fig.suptitle(f"History match: the best settings reproduce the real field to within about {avg:.1f}%",
                 fontsize=15, weight="bold")
    caption(fig, "Red line = best settings found by the proxy search; red band = settings that match nearly as "
                 "well; green dashes = closest of the 100 existing runs.")
    savePlot(fig, "5_history_match.png")


def fromHistoryEnd(fc, s, arr):
    """x, y for a forecast line that starts exactly at the last real history point."""
    return (fc["fdates"].insert(0, fc["histEnd"]),
            np.concatenate([[fc["obsEnd"][s]], np.asarray(arr, float)]))


def drawForecastLayers(ax, fc, s, cmap, label=True, midlw=3.0):
    """Cumulative forecast panel. Draw order (back to front): blue band, the 10 submitted cases, then
    the middle forecast on top with a white outline so the cases can never hide it. The legend names
    every kind of line on the panel."""
    lo, mid, hi = fc["bands"][s]
    (xs, ylo), (_, ymid), (_, yhi) = (fromHistoryEnd(fc, s, a) for a in (lo, mid, hi))
    ax.fill_between(xs, ylo, yhi, color=BLUE, alpha=0.3, zorder=1,
                    label="spread between equally good matches" if label else None)
    for y in (ylo, yhi):
        ax.plot(xs, y, color=BLUE, lw=0.9, ls="--", alpha=0.9, zorder=4)
    curves = fc["selCurves"][s]
    n = len(curves)
    for rank, c in enumerate(curves):
        lab = None
        if label and rank == 0:
            lab = "submitted Case 1 (best history match)"
        elif label and rank == n - 1:
            lab = f"submitted Case {n} (weakest of the {n})"
        ax.plot(*fromHistoryEnd(fc, s, c), color=cmap(rank / max(n - 1, 1)), lw=1.3, alpha=0.95, zorder=3,
                label=lab)
    ax.plot(xs, ymid, color=BLUE, lw=midlw, zorder=5, label="middle forecast" if label else None,
            path_effects=[pe.withStroke(linewidth=midlw + 2.5, foreground="white")])


def plotForecast(data, fc):
    """6: the forecast to 2028 (cumulative, as in the template) + the forecast method tested on the past."""
    fdates, bands, histEnd = fc["fdates"], fc["bands"], fc["histEnd"]
    yr = fdates[-1].year
    fig = plt.figure(figsize=(21, 5.6), constrained_layout=True)
    gs = fig.add_gridspec(1, 4, width_ratios=[1, 1, 1, 0.7])
    cmap = plt.get_cmap("viridis")
    for j, s in enumerate(STREAMS):
        ax = fig.add_subplot(gs[0, j])
        lo, mid, hi = bands[s]
        h = data[f"{s}_cum"][HISTORY_COL].dropna()
        ax.plot(h.index, h.values, "ko", ms=3, label="real field")
        drawForecastLayers(ax, fc, s, cmap, label=(j == 0), midlw=3.0)
        ax.axvline(histEnd, color="k", ls=":")
        ax.text(0.98, 0.05, f"{yr}:  {mid[-1] / 1e6:.1f}M  ({lo[-1] / 1e6:.1f} to {hi[-1] / 1e6:.1f}M)",
                transform=ax.transAxes, ha="right", fontsize=10, weight="bold",
                bbox=dict(fc="white", ec="#999999", alpha=0.9))
        ax.set_title(f"{s.capitalize()} produced so far [{UNITS[s]}]", weight="bold")
        ax.grid(alpha=0.3)
        millions(ax)
        if j == 0:
            ax.legend(fontsize=8.5, loc="upper left")
    e, chosen = fc["errs"][fc["chosen"]], fc["chosen"]
    med = [100 * e[s].abs().median() for s in STREAMS]
    ax = fig.add_subplot(gs[0, 3])
    ax.bar([s.capitalize() for s in STREAMS], med, edgecolor="k",
           color=[GREEN if m < 5 else AMBER if m < 10 else RED for m in med])
    for i, m in enumerate(med):
        ax.text(i, m, f"{m:.1f}%", ha="center", va="bottom", weight="bold")
    ax.set_title(f"Test on the past\n(last {HINDCAST_YEARS:g} yr hidden, then forecast)", weight="bold")
    ax.set_ylabel("typical error of the production increase [%]")
    ax.grid(alpha=0.3, axis="y")
    oil = bands["oil"]
    fig.suptitle(f"Forecast {histEnd.year} to {yr}: about {oil[1][-1] / 1e6:.0f}M STB of oil produced by {yr} "
                 f"(range {oil[0][-1] / 1e6:.0f} to {oil[2][-1] / 1e6:.0f}M)", fontsize=15, weight="bold")
    caption(fig, "Thin coloured lines = the 10 submitted cases (dark purple = best history match, yellow = weakest of the "
                 "10); thick blue line = middle forecast; band = spread between settings that fit history equally "
                 "well.\nRight of the dotted line is extrapolation: real uncertainty is larger.")
    savePlot(fig, "6_forecast_2008_2028.png")


def plotForecastTrajectory(data, fc):
    """7: the full forecast trajectory to 2028 - cumulative (top) and production rate (bottom)."""
    fdates, bands, histEnd, obsEnd = fc["fdates"], fc["bands"], fc["histEnd"], fc["obsEnd"]
    fig, axes = plt.subplots(2, 3, figsize=(19, 10), constrained_layout=True, sharex=True)
    cmap = plt.get_cmap("viridis")
    days = np.diff(np.concatenate([[histEnd.value],
                                   fdates.values.astype("datetime64[ns]").astype("int64")])) / 8.64e13
    for j, s in enumerate(STREAMS):
        ax, ax2 = axes[0, j], axes[1, j]
        lo, mid, hi = bands[s]
        h = data[f"{s}_cum"][HISTORY_COL].dropna()
        ax.plot(h.index, h.values, "ko", ms=3, label="real field (history)")
        drawForecastLayers(ax, fc, s, cmap, label=(j == 0), midlw=2.6)
        hr = data[f"{s}_rate"][HISTORY_COL].dropna()
        ax2.plot(hr.index, hr.values, "ko", ms=3)
        n = len(fc["selCurves"][s])
        for rank, c in enumerate(fc["selCurves"][s]):               # rate of each submitted case
            prev = np.concatenate([[obsEnd[s]], c[:-1]])
            ax2.plot(fdates, (c - prev) / days, color=cmap(rank / max(n - 1, 1)), lw=1.0, alpha=0.85, zorder=3)
        for arr, ls, lw, z in [(lo, "--", 1.0, 4), (hi, "--", 1.0, 4), (mid, "-", 2.6, 5)]:   # blue on top
            prev = np.concatenate([[obsEnd[s]], arr[:-1]])
            kw = dict(path_effects=[pe.withStroke(linewidth=lw + 2.5, foreground="white")]) if ls == "-" else {}
            ax2.plot(fdates, (arr - prev) / days, color=BLUE, ls=ls, lw=lw, zorder=z, **kw)
        for a in (ax, ax2):
            a.axvline(histEnd, color="k", ls=":")
            a.grid(alpha=0.3)
        millions(ax)
        ax.set_title(f"{s.capitalize()} produced so far [{UNITS[s]}]", weight="bold")
        ax2.set_title(f"{s.capitalize()} production rate [{UNITS[s]}/day]", weight="bold")
        ax.text(histEnd, ax.get_ylim()[0], " forecast starts", va="bottom", fontsize=9)
    axes[0, 0].legend(fontsize=8.5, loc="upper left")
    fig.suptitle(f"Forecast {histEnd.year} to {fdates[-1].year}: what the field is expected to produce",
                 fontsize=15, weight="bold")
    fig.supxlabel("Thin coloured lines = the 10 cases submitted (dark purple = best history match, yellow = weakest of "
                  "the 10).  Thick blue line = middle forecast; dashed blue / band = spread between settings that match "
                  "history equally well.\nRight of the dotted line is an EXTRAPOLATION: the real uncertainty is larger "
                  "than the band.", fontsize=11)
    savePlot(fig, "7_forecast_trajectory_2008_2028.png")


# =============================================================================
# STEP 7 - THE TWO RESULT EXCELS: A = history matching, B = 20-year forecast
# =============================================================================
HEAD_FILL = PatternFill("solid", fgColor="305496")
VERDICT_FILL = {"PINNED DOWN": "C6EFCE", "partly constrained": "FFEB9C", "NOT constrained": "FFC7CE"}


def startSheet(wb, title, line):
    ws = wb.create_sheet(title)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = line
    ws["A2"].font = Font(italic=True, color="595959")
    ws.page_setup.orientation = "landscape"                  # prints one page wide
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    return ws


def putTable(ws, df, r0, heading=None, fills=None):
    """Write a DataFrame at row r0 (optional bold heading above it); returns the next free row."""
    if heading:
        ws.cell(row=r0, column=1, value=heading).font = Font(bold=True, size=12)
        r0 += 1
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(row=r0, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start=r0 + 1):
        for j, v in enumerate(row, start=1):
            if isinstance(v, (np.floating, float)):
                v = None if not np.isfinite(v) else float(v)
            elif isinstance(v, np.integer):
                v = int(v)
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            cell = ws.cell(row=i, column=j, value=v)
            if isinstance(v, float):
                cell.number_format = "#,##0" if abs(v) >= 1000 else "0.000"
            elif hasattr(v, "year"):
                cell.number_format = "yyyy-mm-dd"
            if fills and (i - r0 - 1) in fills:
                cell.fill = PatternFill("solid", fgColor=fills[i - r0 - 1])
    for j, c in enumerate(df.columns, start=1):
        letter = get_column_letter(j)
        width = max([len(str(c)) * 0.8] + [len(f"{v:,.0f}" if isinstance(v, (int, float)) else str(v))
                                           for v in df[c].head(60)])
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0,
                                                 min(max(width + 3, 11), 44))
    return r0 + len(df) + 2


def writeSummarySheet(wb, line, rows):
    """A 'Summary' sheet: Question | Answer | Where to look."""
    ws = startSheet(wb, "Summary", line)
    for j, h in enumerate(["Question", "Answer", "Where to look"], start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
    for i, (q, a, w) in enumerate(rows, start=5):
        for j, v in enumerate((q, a, w), start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=1).font = Font(bold=True)
        ws.row_dimensions[i].height = 15 * (1 + len(a) // 95)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 38, 100, 30


def writeHistoryMatchWorkbook(hm, rank):
    """PART A excel: only history matching (which cases / settings fit the real field)."""
    wb = Workbook()
    wb.remove(wb.active)
    pcols, rep, ident = hm["pcols"], hm["rep"], hm["ident"]
    held = rep[rep["dataset"].isin(["validation", "test"])]
    held = held if len(held) else rep
    err, worst = held["typical_error_pct"].mean(), held["worst_error_pct"].max()
    grp = lambda v: ", ".join(ident.loc[ident["verdict"] == v, "setting"]) or "none"
    bestEx = rank.index[0]
    best = "   |   ".join(f"{c} {v:.4g}" for c, v in zip(pcols, hm["bestParams"]))

    # ---- Summary ----
    writeSummarySheet(wb, "PART A - HISTORY MATCHING. The answers in one page. Graphs 1-5 (PNG files) show the same "
                          "results visually. The 20-year forecast is in a separate file.", [
        ("Which settings reproduce the real history?", f"{best}   (misfit {hm['bestMisfit']:.3f})",
         "Graph 5 | sheet Settings"),
        ("Which existing case is closest?", f"{bestEx} (misfit {rank['misfit'].iloc[0]:.3f}). The proxy-found "
         f"settings reach {hm['bestMisfit']:.3f} but are not a simulator run.", "Graph 3 | sheet Case ranking"),
        ("What does the history pin down?", f"Pinned down: {grp('PINNED DOWN')}.  Partly: {grp('partly constrained')}.  "
         f"Not determined: {grp('NOT constrained')}.", "Graph 4 | sheet Settings"),
        ("Can the proxy be trusted?", f"{verdictWord(err).capitalize()}: typical error {err:.1f}% (worst case "
         f"{worst:.1f}%) on simulated cases it never saw.", "Graph 2 | sheet Proxy accuracy"),
        ("Method", "Proxy = PCA + Gaussian process (suits ~85 runs and 4 settings). Misfit = oil and water cumulative "
         "plus water rate error against the real history.", ""),
        ("Train / validation / test", f"The split of the cases is {SPLIT_SOURCE}. It only tests the proxy; it plays "
         "no role in the ranking of cases.", "sheet Data split"),
    ])

    # ---- Case ranking (one score, all existing cases, proxy optimum on top) ----
    ws = startSheet(wb, "Case ranking", "All simulated cases ranked by misfit to the real history (lower = closer). "
                    "Green = top 10. Red row = best settings found by the proxy (not a simulator run).")
    top = pd.DataFrame({"Rank": ["-"], "Case": ["Best settings found (proxy)"],
                        "Misfit": [hm["bestMisfit"]], **{c: [v] for c, v in zip(pcols, hm["bestParams"])},
                        **{SHORT[k]: [100 * v] for k, v in hm["resid"].items()}})
    body = pd.DataFrame({"Rank": rank["rank"].values, "Case": rank.index,
                         "Misfit": rank["misfit"].values, **{c: rank[c].values for c in pcols},
                         **{SHORT[k]: rank[f"err_{k}"].values for k in hm["targets"]}})
    allRows = pd.concat([top, body], ignore_index=True)
    fills = {0: "FFC7CE", **{i: "E2EFDA" for i in range(1, 11)}}
    putTable(ws, allRows, 4, fills=fills)
    ws.auto_filter.ref = f"A4:{get_column_letter(len(allRows.columns))}{4 + len(allRows)}"
    ws.freeze_panes = "C5"

    # ---- Settings ----
    ws = startSheet(wb, "Settings", "The 4 uncertain settings: best value found and how far the real history determines each.")
    meaning = {"PINNED DOWN": "History determines it: only a narrow range matches.",
               "partly constrained": "History narrows it somewhat.",
               "NOT constrained": "History cannot tell: any value in the range fits."}
    st = ident.copy()
    st["share_of_range"] = 100 * st["share_of_range"]
    st["meaning"] = st["verdict"].map(meaning)
    st.columns = ["Setting", "Best value", "Plausible low", "Plausible high", "Tested low", "Tested high",
                  "Plausible share of tested range [%]", "Verdict", "What it means"]
    putTable(ws, st, 4, fills=None)
    for i, v in enumerate(st["Verdict"], start=5):
        ws.cell(row=i, column=8).fill = PatternFill("solid", fgColor=VERDICT_FILL[v])
        for j in range(2, 7):
            ws.cell(row=i, column=j).number_format = "0.000"

    # ---- Proxy accuracy (this is where train / validation / test belongs) ----
    ws = startSheet(wb, "Proxy accuracy", "How far to trust the fast ML stand-in (proxy). Train = cross-validation; "
                    "validation / test = cases the proxy did not train on.")
    r = rep.copy()
    r["measure"] = r["measure"].map(lambda k: PLAIN.get(k, k))
    r.columns = ["Dataset", "Predicted", "Cases", "Typical error [% of largest history value]",
                 "Worst error [%]", "R2 at end of history", "Verdict"]
    putTable(ws, r, 4)

    # ---- Data split (where train / validation / test comes from) ----
    ws = startSheet(wb, "Data split", f"Train / validation / test: {SPLIT_SOURCE}.")
    r0 = 4
    if len(SPLIT_REPORT):
        r0 = putTable(ws, SPLIT_REPORT.rename(columns={"file": "File", "split": "Used as", "cases found": "Cases found",
                                                       "found via": "Case numbers read from", "first ids": "First ids"}),
                      4, heading="What was found in files 03-08")
    sp = pd.DataFrame({"Case": hm["cases"], "Used as": [splitOf(c) for c in hm["cases"]]})
    putTable(ws, sp, r0, heading="Split of every case used")

    path = OUTPUT_DIR / HM_XLSX
    wb.save(path)
    print(f"saved {path}")


def writeForecastWorkbook(fc):
    """PART B excel: only the 20-year forecast and its reliability."""
    wb = Workbook()
    wb.remove(wb.active)
    yr = fc["fdates"][-1].year
    bands, e = fc["bands"], fc["errs"][fc["chosen"]]
    fcText = "   |   ".join(f"{s.capitalize()} {bands[s][1][-1]:,.0f} ({bands[s][0][-1]:,.0f} to {bands[s][2][-1]:,.0f}) "
                            f"{UNITS[s]}" for s in STREAMS)
    npass = int(fc["check"]["result"].str.startswith("PASS").sum())

    # ---- Summary ----
    writeSummarySheet(wb, f"PART B - {fc['histEnd'].year} TO {yr} FORECAST. The answers in one page. Graphs 6-7 "
                          "(PNG files) show the same results visually. History matching is in a separate file.", [
        (f"Forecast: produced by {yr}", f"{fcText}   (middle, with 10th to 90th percentile)",
         "Graphs 6-7 | sheet Forecast"),
        ("Which cases were submitted?", f"{len(fc['selected'])} cases: Case 1 matches history best, the rest differ "
         "on purpose to show the plausible range.", "sheet Submitted cases"),
        ("How reliable is the forecast method?", f"Tested on the past (last {HINDCAST_YEARS:g} years hidden): typical "
         "error " + ", ".join(f"{s} {100 * e[s].abs().median():.1f}%" for s in STREAMS)
         + f". A {yr - fc['histEnd'].year}-year forecast will be less accurate.", "Graph 6 | sheet Forecast accuracy"),
        ("Is the template filled correctly?", f"{npass} of {len(fc['check'])} sheets pass (no blanks, no negatives, "
         "cumulative never decreases).", "sheet Forecast accuracy"),
        ("Method", "Forecast = decline-curve extrapolation per case, then a Gaussian-process proxy across settings. "
         "RNN / autoencoder / DMD were not used: too little data to train them.", ""),
        ("Limits", f"{fc['histEnd'].year}-{yr} is extrapolation (data and simulator stop in {fc['histEnd'].year}); the "
         "injection / well plan after that is not in the data.", ""),
    ])

    # ---- Submitted cases ----
    ws = startSheet(wb, "Submitted cases", "The cases in the template. Case 1 matches history best; the rest are chosen to "
                    "differ, so together they show the plausible range. They are new settings, not existing runs.")
    sel = fc["selected"].copy()
    sel["case"] = [f"Case {k}" for k in sel["case"]]
    sel = sel.rename(columns={"case": "Submitted as", "misfit": "Misfit (proxy)",
                              f"gas_cum_{yr}": f"Gas by {yr} [MSCF]", f"oil_cum_{yr}": f"Oil by {yr} [STB]",
                              f"water_cum_{yr}": f"Water by {yr} [STB]"})
    putTable(ws, sel, 4)

    # ---- Forecast ----
    ws = startSheet(wb, "Forecast", f"Cumulative production per template date: middle forecast and 10th / 90th percentile "
                    f"across the {fc['nAcceptable']} settings that match history about equally well.")
    band = {"Date": fc["fdates"]}
    for s in STREAMS:
        for nm, arr in zip(["10th pct", "middle", "90th pct"], bands[s]):
            band[f"{s.capitalize()} {nm} [{UNITS[s]}]"] = arr
    putTable(ws, pd.DataFrame(band), 4)
    ws.freeze_panes = "B5"

    # ---- Forecast accuracy ----
    ws = startSheet(wb, "Forecast accuracy", "How far to trust the forecast method and the filled template.")
    hind = pd.DataFrame([{"Fitting window [years]": w, "Chosen": "yes" if w == fc["chosen"] else "",
                          **{f"{s.capitalize()} error [%]": 100 * ee[s].abs().median() for s in STREAMS}}
                         for w, ee in fc["errs"].items()])
    r0 = putTable(ws, hind, 4, heading=f"Forecast method tested on the past (last {HINDCAST_YEARS:g} years hidden)")
    putTable(ws, fc["check"].rename(columns={"sheet": "Template sheet", "rows": "Rows", "result": "Check"}),
             r0, heading="Filled template check (read back from the saved file)")

    path = OUTPUT_DIR / FC_XLSX
    wb.save(path)
    print(f"saved {path}")


# =============================================================================
# MAIN
# =============================================================================
def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print(f"Reading from: {DATA_DIR.resolve()}")
    params = loadParameters()
    data = {key: loadProduction(sheet) for key, sheet in PRODUCTION_SHEETS.items()}
    buildSplitMap(caseCols(data["oil_cum"]))
    printDataChecks(data, params)

    hm = runHistoryMatch(data, params)
    rank = rankCases(params, hm)
    plotProblem(hm)
    plotProxyTrust(hm)
    plotRanking(hm, rank)
    plotWhatHistoryTells(hm)
    plotHistoryMatch(hm)
    writeHistoryMatchWorkbook(hm, rank)              # PART A: history matching only

    fc = runForecast(data, params, hm)               # PART B: 20-year forecast only
    plotForecast(data, fc)
    plotForecastTrajectory(data, fc)
    writeForecastWorkbook(fc)
    print("\nREMINDERS: (1) 2008-2028 is an extrapolation; the injection / well plan after 2008 is not in "
          "the data.\n(2) Confirm what 'Case 1..10' mean and how they are scored; change SELECTION_MODE if "
          "needed.")
    if SHOW_PLOTS:
        plt.show()


if __name__ == "__main__":
    main()